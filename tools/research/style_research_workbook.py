#!/usr/bin/env python3
"""Format a verified research workbook without changing cell values or formulas.
Keeps existing formula caches, comments, links, and workbook features; writes a copy.
Run with the Workspace .venv interpreter, supplying Workspace-relative file paths.
"""
import argparse,copy,hashlib,json,math,re,unicodedata
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from openpyxl.styles import Alignment,Border,Font,PatternFill,Side
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.views import Selection
from openpyxl.utils import get_column_letter
NAVY='18324F';TEXT='243447';MUTED='64748B';BLUE='2454A4';TEAL='20746C';LINE='CBD5E1';PAPER='FFFFFF'
LIGHT='F2F5F9';TOTAL='EAF0F8';CHECK='EDF5F1';AMBER='806741';QHDR='DDEEEB'
NUM='#,##0;(#,##0);"–"';NUM1='#,##0.0;(#,##0.0);"–"';PCT='0.0%;(0.0%);"–"';PS='0.00;(0.00);"–"';MULT='0.0"x";(0.0"x");"–"'
NS={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}

def snapshot(w):
 return {(s.title,c.coordinate):(c.value,c.data_type) for s in w for row in s for c in row if c.value is not None}

def merge_empty(ws,row,start,end):
 if end<=start:return
 if any(ws.cell(row,c).value is not None for c in range(start+1,end+1)):return
 for r in ws.merged_cells.ranges:
  if r.min_row<=row<=r.max_row:return
 ws.merge_cells(start_row=row,start_column=start,end_row=row,end_column=end)

def text_of(c):
 s=str(c.value or '')
 if c.data_type=='f' and s.startswith('=HYPERLINK'):
  m=re.search(r',"(.+)"\)$',s)
  return m[1] if m else s
 return s

def height(ws,rn,visible_cols=None,min_height=21):
 nlines=1
 for c in ws[rn]:
  if c.value is None or (visible_cols is not None and c.column not in visible_cols):continue
  if isinstance(c.value,(int,float)) or c.data_type=='f' and not c.value.startswith('=HYPERLINK'):continue
  width=ws.column_dimensions[c.column_letter].width or 13
  if ws.column_dimensions[c.column_letter].hidden:continue
  for mr in ws.merged_cells.ranges:
   if mr.min_row==rn and mr.min_col==c.column:
    width=sum((ws.column_dimensions[get_column_letter(i)].width or 13) for i in range(mr.min_col,mr.max_col+1) if not ws.column_dimensions[get_column_letter(i)].hidden)
    break
  units=sum(2 if unicodedata.east_asian_width(ch) in 'WF' else 1 for ch in text_of(c))
  usable=max(8,width-4-(c.alignment.indent or 0)*3)
  sz=c.font.sz or 10.5
  nlines=max(nlines,math.ceil(units*sz/10.5/usable))
 return max(min_height,round(nlines*15.5+6,1))

def section(ws,rn,end,merged=True):
 if merged:merge_empty(ws,rn,1,end)
 for col in range(1,end+1):
  c=ws.cell(rn,col);c.fill=PatternFill('solid',fgColor=NAVY);c.border=Border()
  c.font=Font(name='Arial',size=11,bold=True,color=PAPER)
  c.alignment=Alignment(vertical='center',wrap_text=True)
 ws.row_dimensions[rn].height=height(ws,rn,min_height=29)

def base(ws,last,statement=False):
 ws.sheet_view.showGridLines=False;ws.sheet_view.zoomScale=95 if statement else 100
 ws.sheet_view.zoomScaleNormal=ws.sheet_view.zoomScale
 ws.sheet_format.defaultRowHeight=21
 for row in ws:
  for c in row:
   if c.value is None:continue
   c.font=Font(name='Arial',size=10.5,color=TEXT)
   c.alignment=Alignment(vertical='center',wrap_text=isinstance(c.value,str) and c.data_type!='f')
   c.border=Border();c.fill=PatternFill(fill_type=None)
 if ws.title!='Audit 2026-10-14':
  section(ws,1,last);ws['A1'].font=Font(name='Arial',size=16,bold=True,color=PAPER);ws.row_dimensions[1].height=36
  merge_empty(ws,2,1,last);ws['A2'].font=Font(name='Arial',size=9,color=MUTED);ws['A2'].alignment=Alignment(vertical='center');ws.row_dimensions[2].height=21
 ws.sheet_properties.pageSetUpPr.fitToPage=True
 ws.page_setup.orientation='landscape' if statement or ws.title!='Cover' else 'portrait'
 ws.page_setup.paperSize=ws.PAPERSIZE_A3 if statement else ws.PAPERSIZE_A4
 ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
 ws.page_margins.left=.25;ws.page_margins.right=.25;ws.page_margins.top=.4;ws.page_margins.bottom=.4
 ws.oddFooter.left.text='KO | Company research';ws.oddFooter.left.size=8;ws.oddFooter.left.font='Arial';ws.oddFooter.left.color=MUTED
 ws.oddFooter.right.text='Page &P / &N';ws.oddFooter.right.size=8;ws.oddFooter.right.color=MUTED
 ws.print_options.horizontalCentered=False


def statements(ws,original):
 old=original[ws.title];headlines={r[0].row for r in old if r[0].font.bold}
 old_sections={r[0].row for r in old if r[0].value and (r[0].fill.fgColor.rgb or '')[-6:]=='1F4E79'}
 ws.column_dimensions['A'].width=64
 ws.column_dimensions['B'].width=34;ws.column_dimensions['B'].hidden=True;ws.column_dimensions['B'].outlineLevel=1
 for col in range(3,22):ws.column_dimensions[get_column_letter(col)].width=11.6
 for col in [13,14,15]:ws.column_dimensions[get_column_letter(col)].width=2.5
 # Remove imported grouped-column width ranges so Excel sees explicit widths.
 for key in list(ws.column_dimensions):
  if key not in [get_column_letter(c) for c in range(1,22)]:del ws.column_dimensions[key]
 base(ws,21,True)
 ws.freeze_panes='C6';ws.print_title_rows='1:4';ws.print_area=f'A1:U{ws.max_row}'
 for rn in range(3,ws.max_row+1):
  a=ws.cell(rn,1);label=text_of(a);row=ws[rn]
  hasdata=any(c.value is not None for c in row[2:12]+row[15:21])
  if rn in [3,4]:
   for col in list(range(3,13))+list(range(16,22)):
    c=ws.cell(rn,col);c.fill=PatternFill('solid',fgColor=NAVY if col<13 and rn==3 else TEAL if rn==3 else LIGHT)
    c.font=Font(name='Arial',size=10 if rn==3 else 8.5,bold=rn==3,color=PAPER if rn==3 else MUTED)
    c.alignment=Alignment(horizontal='center',vertical='center');c.border=Border(bottom=Side(style='thin',color=LINE)) if rn==4 else Border()
   for col in [13,14,15]:ws.cell(rn,col).font=Font(name='Arial',size=8,color='B5C0CB');ws.cell(rn,col).fill=PatternFill('solid',fgColor='F8FAFC')
   ws.row_dimensions[rn].height=25 if rn==3 else 18;continue
  if not label and not hasdata:
   ws.row_dimensions[rn].height=10;continue
  if a.data_type=='f' and a.value.startswith('=HYPERLINK'):
   merge_empty(ws,rn,1,21);a.font=Font(name='Arial',size=9.5,color=TEAL,underline='single');a.alignment=Alignment(vertical='center',wrap_text=True)
   ws.row_dimensions[rn].height=height(ws,rn,min_height=22);continue
  if rn in old_sections or (not hasdata and (re.match(r'^\d+\. ',label) or rn==123 and ws.title=='Operating Metrics' or rn==104 and ws.title=='Note Schedules')):
   section(ws,rn,21);continue
  if label=='CONTENTS':
   merge_empty(ws,rn,1,21);a.font=Font(name='Arial',size=9,bold=True,color=MUTED);ws.row_dimensions[rn].height=22;continue
  ischeck=label.lower().startswith(('check:','balance check','cash walk check','cash walk check:'))
  isinfo=label.startswith('Info:')
  issub=bool(label) and not hasdata
  bold=rn in headlines and not ischeck and not isinfo
  # Section-like supplemental totals retain the same hierarchy as the main statements.
  if ws.title=='Note Schedules' and rn in [123,124,125]:bold=True
  a.font=Font(name='Arial',size=10.5,bold=bold or issub,color=MUTED if isinfo else TEXT,italic=isinfo or label.startswith('  '))
  a.alignment=Alignment(vertical='center',wrap_text=True,indent=old.cell(rn,1).alignment.indent or 0)
  if ws.cell(rn,2).value is not None:
   ws.cell(rn,2).font=Font(name='Arial',size=9,color=MUTED);ws.cell(rn,2).alignment=Alignment(wrap_text=True,vertical='top')
  bg=CHECK if ischeck else TOTAL if bold and hasdata else LIGHT if issub else None
  for col in list(range(1,13))+list(range(16,22)):
   c=ws.cell(rn,col)
   if bg:c.fill=PatternFill('solid',fgColor=bg)
   if bold and hasdata:c.border=Border(top=Side(style='thin',color=LINE))
  for c in row[2:12]+row[15:21]:
   if c.value is None:continue
   if c.data_type=='f':color=BLUE if re.fullmatch(r'=[0-9.+-]+',c.value) else TEAL if '!' in c.value else TEXT
   else:color=AMBER if c.value=='N/A' else BLUE
   c.font=Font(name='Arial',size=10.5,color=MUTED if isinfo else color,bold=bold)
   c.alignment=Alignment(horizontal='center' if c.value=='N/A' else 'right',vertical='center')
   fmt=old[c.coordinate].number_format
   if '%' in fmt:c.number_format=PCT
   elif '"x"' in fmt or '"×"' in fmt:c.number_format=MULT
   elif '0.00' in fmt:c.number_format=PS
   elif '0.0' in fmt:c.number_format=NUM1
   elif c.value!='N/A':c.number_format=NUM
  ws.row_dimensions[rn].height=height(ws,rn)
 ws.sheet_view.selection=[Selection(pane='topRight',activeCell='C3',sqref='C3'),Selection(pane='bottomLeft',activeCell='A6',sqref='A6'),Selection(pane='bottomRight',activeCell='C6',sqref='C6')]
 ws.sheet_properties.tabColor=NAVY if ws.title in ['Income Statement','Balance Sheet','Cash Flow'] else '507D79' if ws.title in ['Segments','Geography'] else '706487' if ws.title=='Derived Metrics' else 'A3824F'


def cover(ws,old):
 ws.column_dimensions['A'].width=34;ws.column_dimensions['B'].width=108
 for col in range(3,9):ws.column_dimensions[get_column_letter(col)].width=2
 base(ws,2);ws.freeze_panes='B6';ws.print_area=f'A1:B{ws.max_row}'
 for row in ws.iter_rows(min_row=3):
  rn=row[0].row;a,b=row[:2]
  if all(c.value is None for c in row):ws.row_dimensions[rn].height=10;continue
  if a.value is not None and b.value is None:
   section(ws,rn,2);continue
  a.font=Font(name='Arial',size=10.5,bold=True,color=TEXT);a.alignment=Alignment(wrap_text=True,vertical='top')
  b.font=Font(name='Arial',size=10.5,color=TEXT);b.alignment=Alignment(wrap_text=True,vertical='top')
  for c in [a,b]:c.fill=PatternFill('solid',fgColor='F7F9FC' if rn%2==0 else PAPER)
  if a.value in ['Blue font','Black font','Green font','Grey italic']:
   a.font=Font(name='Arial',size=10.5,bold=True,color={'Blue font':BLUE,'Black font':TEXT,'Green font':TEAL,'Grey italic':MUTED}[a.value])
  ws.row_dimensions[rn].height=height(ws,rn,min_height=24)
 ws.sheet_properties.tabColor=NAVY


def sources(ws,old):
 for col,width in zip('ABCDE',[25,36,84,45,43]):ws.column_dimensions[col].width=width
 for col in 'FGH':ws.column_dimensions[col].width=2
 base(ws,5);ws.freeze_panes='C5';ws.print_title_rows='1:4';ws.print_area=f'A1:E{ws.max_row}'
 for row in ws.iter_rows(min_row=3):
  rn=row[0].row;nonnull=[c for c in row if c.value is not None]
  if not nonnull:ws.row_dimensions[rn].height=10;continue
  if len(nonnull)==1 and nonnull[0].column==1:section(ws,rn,5);continue
  if rn in [4,23,44,57]:
   for c in row[:5]:c.fill=PatternFill('solid',fgColor=NAVY);c.font=Font(name='Arial',size=10,bold=True,color=PAPER);c.alignment=Alignment(wrap_text=True,vertical='center')
  else:
   if len(nonnull)<=2 and all(c.value is None for c in row[2:5]):merge_empty(ws,rn,2,5)
   for c in row[:5]:
    c.font=Font(name='Arial',size=9.5 if rn<=20 or 58<=rn<=63 else 10.5,color=TEXT,bold=c.column==1)
    c.alignment=Alignment(wrap_text=True,vertical='top',horizontal='right' if isinstance(c.value,(int,float)) else 'left')
    c.fill=PatternFill('solid',fgColor='F7F9FC' if rn%2==0 else PAPER)
    if isinstance(c.value,(int,float)):c.number_format=NUM
  ws.row_dimensions[rn].height=height(ws,rn,visible_cols=range(1,6),min_height=25)
 ws.sheet_properties.tabColor='718096'


def audit(ws,old):
 for col,width in zip('ABCDEFGHIJ',[26,12,15,58,18,30,24,60,75,64]):ws.column_dimensions[col].width=width
 base(ws,7);ws.freeze_panes='E23';ws.print_title_rows='22:22';ws.print_area=f'A1:J{ws.max_row}'
 ws.column_dimensions.group('H','J',outline_level=1,hidden=True);ws.sheet_properties.outlinePr.summaryRight=False
 section(ws,1,7);ws['A1'].font=Font(name='Arial',size=16,bold=True,color=PAPER);ws.row_dimensions[1].height=36
 for rn in range(2,ws.max_row+1):
  if rn in [12,21]:ws.row_dimensions[rn].height=10;continue
  if rn<=11 and isinstance(ws.cell(rn,2).value,str):merge_empty(ws,rn,2,7)
  if rn in [13,22]:
   for c in ws[rn]:c.fill=PatternFill('solid',fgColor=NAVY);c.font=Font(name='Arial',size=10,bold=True,color=PAPER);c.alignment=Alignment(vertical='center',wrap_text=True)
  else:
   for c in ws[rn]:
    c.font=Font(name='Arial',size=10.5,color=TEXT,bold=rn<=11 and c.column==1)
    c.alignment=Alignment(wrap_text=True,vertical='top',horizontal='right' if isinstance(c.value,(int,float)) else 'left')
    c.fill=PatternFill('solid',fgColor='F7F9FC' if rn%2==0 else PAPER)
  ws.row_dimensions[rn].height=height(ws,rn,visible_cols=range(1,8),min_height=25)
 ws.sheet_properties.tabColor='718096'


def preserve_cache(src,dst):
 with ZipFile(src) as z:before={n:z.read(n) for n in z.namelist()}
 with ZipFile(dst) as z:parts={n:z.read(n) for n in z.namelist()}
 count=0
 for n,b in list(parts.items()):
  if not n.startswith('xl/worksheets/sheet') or not n.endswith('.xml') or n not in before:continue
  old=ET.fromstring(before[n]);new=ET.fromstring(b)
  cells={c.get('r'):c for c in old.findall('.//s:c',NS) if c.find('s:f',NS) is not None}
  for c in new.findall('.//s:c',NS):
   f=c.find('s:f',NS);o=cells.get(c.get('r'))
   if f is None or o is None or f.text!=o.find('s:f',NS).text:continue
   v=o.find('s:v',NS)
   if v is None:continue
   current=c.find('s:v',NS)
   if current is not None:c.remove(current)
   c.append(copy.deepcopy(v))
   if o.get('t'):c.set('t',o.get('t'))
   else:c.attrib.pop('t',None)
   count+=1
  parts[n]=ET.tostring(new,encoding='utf-8',xml_declaration=True)
 with ZipFile(dst,'w',ZIP_DEFLATED) as z:
  for n,b in parts.items():z.writestr(n,b)
 return count

def main():
 p=argparse.ArgumentParser();p.add_argument('source');p.add_argument('output');p.add_argument('--report');a=p.parse_args();src=Path(a.source);dst=Path(a.output)
 if dst.exists():raise FileExistsError('Preserve existing output: '+str(dst))
 before=load_workbook(src);w=load_workbook(src);data=snapshot(before)
 for ws in w:
  if ws.title=='Cover':cover(ws,before[ws.title])
  elif ws.title=='Sources':sources(ws,before[ws.title])
  elif ws.title.startswith('Audit '):audit(ws,before[ws.title])
  else:statements(ws,before)
 assert snapshot(w)==data,'A value or formula changed during formatting'
 for ws in w:ws.sheet_view.tabSelected=ws.title=='Operating Metrics'
 w.active=w.sheetnames.index('Operating Metrics');w['Operating Metrics'].sheet_view.selection[-1]=Selection(pane='bottomRight',activeCell='C14',sqref='C14')
 dst.parent.mkdir(parents=True,exist_ok=True);w.save(dst);n=preserve_cache(src,dst)
 actual=load_workbook(dst);assert snapshot(actual)==data
 report={'source':str(src),'output':str(dst),'source_sha256':hashlib.sha256(src.read_bytes()).hexdigest(),'output_sha256':hashlib.sha256(dst.read_bytes()).hexdigest(),'cell_values_and_formulas_unchanged':True,'formula_caches_preserved':n,'sheets':[]}
 for ws in actual:
  fonts=sorted({c.font.name for row in ws for c in row if c.value is not None})
  report['sheets'].append({'name':ws.title,'font_families':fonts,'zoom':ws.sheet_view.zoomScale,'rows':ws.max_row,'merged_ranges':len(ws.merged_cells.ranges),'hidden_provenance':ws.title.startswith('Audit '),'row_height_min':min(d.height for d in ws.row_dimensions.values() if d.height),'row_height_max':max(d.height for d in ws.row_dimensions.values() if d.height)})
 if a.report:Path(a.report).write_text(json.dumps(report,ensure_ascii=False,indent=2))
 print(json.dumps(report,ensure_ascii=False))
if __name__=='__main__':main()
