#!/usr/bin/env python3
"""Verify declarations, reopened originals, model arithmetic and EVERY workbook cell.

Exit 0: checks and bound independent-review record pass; 2: failure; 3: unable to verify.
Distinct declared Agent IDs cannot prove an independent review actually occurred.
A Python recomputation never counts as an Excel-compatible engine calculation.
Visual inspection is separately required by the skill before release.
"""
from __future__ import annotations

import argparse
import json
import math
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from model_contract import is_number, load_style, model_issues, sanitize_sheet_name, ordered_sheets, tab_color, display_value
from model_engine import Evaluator, FormulaPlan
from source_audit import audit_sources
from agent_review import verify_agent_review

ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}


def rgb(color):
    return color.rgb[-6:] if color is not None and color.type == "rgb" and isinstance(color.rgb, str) else None


def verify(model, workbook, workspace, agent_review=None):
    findings, counts = [], {"input_cells": 0, "formula_cells": 0, "check_cells": 0}
    def report(status, location, message):
        findings.append({"status": status, "location": location, "message": str(message)})
    try:
        for issue in model_issues(model): report("FAIL", "contract", issue)
        source_audit = audit_sources(model, workspace)
        findings.extend(source_audit["findings"])
        plan, ev = FormulaPlan(model), Evaluator(model)
    except Exception as exc:
        report("FAIL", "model/source contract", exc)
        return {"status": "FAIL", "findings": findings, "counts": counts}
    try:
        formulas = load_workbook(workbook, data_only=False)
        cached = load_workbook(workbook, data_only=True)
    except Exception as exc:
        report("FAIL", "workbook", exc)
        return {"status": "FAIL", "findings": findings, "counts": counts}
    style = load_style()
    expected_sheets = {sanitize_sheet_name(s["name"]) for s in model["sheets"]}
    if set(formulas.sheetnames) != expected_sheets:
        report("FAIL", "sheets", "Workbook sheet set differs from the model")
    if formulas.sheetnames != [sanitize_sheet_name(s['name']) for s in ordered_sheets(model['sheets'])]:
        report('FAIL','sheets','Expected Contents, operating schedules, IS, BS, CF, metrics, then notes')
    for sheet in model["sheets"]:
        name = sanitize_sheet_name(sheet["name"])
        if name not in formulas:
            report("FAIL", name, "Worksheet missing")
            continue
        ws, cw = formulas[name], cached[name]
        if rgb(ws.sheet_properties.tabColor) != tab_color(sheet):
            report('FAIL',name,'Tab color differs from its functional group / explicit override')
        if ws.sheet_view.showGridLines is not False:
            report("FAIL", name, "Gridlines must be hidden")
        title = ws["A1"]
        if (title.font.name != style["font"] or title.font.sz != style["sizes"]["title"]
                or not title.font.bold or rgb(title.font.color) != style["colors"]["white"]
                or rgb(title.fill.fgColor) != style["colors"]["title"]):
            report("FAIL", f"{name}!A1", "Title differs from approved Arial/size/white-on-blue style")
        if ws.row_dimensions[1].height is None or ws.row_dimensions[1].height < style["heights"]["title"]:
            report("FAIL", name, "Title row is too short")
        if sheet.get("kind") == "table":
            for ri, record in enumerate(sheet.get("table", []), 4):
                for ci, value in enumerate(record, 1):
                    if ws.cell(ri, ci).value != display_value(value):
                        report("FAIL", f"{name}!{get_column_letter(ci)}{ri}", "Schedule/metadata cell differs from source-bound model")
                    if is_number(value):
                        counts["input_cells"] += 1
                        if not is_number(ws.cell(ri, ci).value):
                            report("FAIL", f"{name}!{get_column_letter(ci)}{ri}", "Reported schedule value lost its numeric type")
            continue
        if sheet.get("kind") != "statement": continue
        short = sheet["short"]
        if ws.freeze_panes != "C6": report("FAIL", name, "Unexpected frozen panes")
        for label, ci in plan.layout.col_of.items():
            if ws.cell(3, ci).value != label:
                report("FAIL", name, f"Period header mismatch at column {ci}: {label}")
            header = ws.cell(3, ci)
            color = "annual_header" if label in plan.layout.annual_labels else "quarter_header"
            if (rgb(header.font.color) != style["colors"]["white"]
                    or rgb(header.fill.fgColor) != style["colors"][color]
                    or header.font.name != style["font"] or header.font.sz != style["sizes"]["header"]):
                report("FAIL", name, f"Period header style differs from profile: {label}")
            end = ws.cell(4, ci).value
            expected_end = plan.layout.info[label].get("end")
            if end is None or not hasattr(end, "isoformat") or end.isoformat()[:10] != expected_end:
                report("FAIL", name, f"Period end is not the expected real date: {label}")
        for bi, block in enumerate(sheet.get("blocks", [])):
            band = ws.cell(plan.block_rows[short][bi], 1)
            if (rgb(band.font.color) != style["colors"]["white"]
                    or rgb(band.fill.fgColor) != style["colors"]["title"]
                    or band.font.name != style["font"] or band.font.sz != style["sizes"]["section"]
                    or not band.font.bold):
                report("FAIL", f"{name}!{band.coordinate}", "Section title style differs from profile")
            for row in block.get("rows", []):
                if row.get("type") in {"spacer", "note"}: continue
                rid, r = row["id"], plan.registry[short][row["id"]]
                expected_label = row.get("label", rid)
                if ws.cell(r, 1).value != expected_label:
                    report("FAIL", f"{name}!A{r}", "Metric label differs from model")
                label_font = ws.cell(r, 1).font
                if label_font.name != style["font"] or label_font.sz != style["sizes"]["body"]:
                    report("FAIL", f"{name}!A{r}", "Body label font differs from profile")
                total = row.get("bold") or row.get("type") == "subtotal"
                ratio = row.get("fmt") in {"pct", "x"} or row.get("type", "").startswith("growth_")
                minimum = style["heights"]["total" if total else "ratio" if ratio else "body"]
                if ws.row_dimensions[r].height is None or ws.row_dimensions[r].height < minimum:
                    report("FAIL", f"{name}!A{r}", "Body row is below approved minimum height")
                for label, ci in plan.layout.col_of.items():
                    addr = f"{name}!{get_column_letter(ci)}{r}"
                    try:
                        expected = plan.cell_value(short, row, label)
                        actual, result = ws.cell(r, ci).value, cw.cell(r, ci).value
                        if actual != expected:
                            report("FAIL", addr, f"Cell differs from model; expected {expected!r}, found {actual!r}")
                            continue
                        if not plan.layout.in_span(row, label): continue
                        if is_number(expected):
                            counts["input_cells"] += 1
                            if not is_number(actual):
                                report("FAIL", addr, "Reported numeric cell lost its numeric type")
                        elif isinstance(expected, str) and expected.startswith("="):
                            counts["formula_cells"] += 1
                            arithmetic = ev.value(short, rid, label)
                            if result in ERRORS:
                                report("FAIL", addr, f"Excel engine error: {result}")
                            elif is_number(arithmetic):
                                if result is None:
                                    report("UNVERIFIABLE", addr, "Missing numeric engine cache; an intentional missing-data blank is different")
                                elif not is_number(result) or not math.isclose(result, arithmetic, rel_tol=1e-10, abs_tol=1e-9):
                                    report("FAIL", addr, f"Engine result {result!r} differs from independent arithmetic {arithmetic}")
                            elif result not in {None, "", "n.m.", "not comparable"}:
                                report("FAIL", addr, "Engine returned a value despite unavailable model inputs")
                            else:
                                negative_base = False
                                if row.get("type", "").startswith("growth_"):
                                    prior = plan.layout.prior(label, "yoy" if row["type"] == "growth_yoy" else "qoq")
                                    if prior:
                                        a = ev.value(*plan.target(row["of"], short, label))
                                        b = ev.value(*plan.target(row["of"], short, prior))
                                        negative_base = is_number(a) and is_number(b) and b <= 0
                                if row.get("type") == "check" or not (negative_base and result == "n.m.") and label not in row.get("unavailable_reason_by_period", {}):
                                    report("UNVERIFIABLE", addr, "Formula depends on missing/noncomparable data; explain the unavailable period")
                        elif expected in {None, "", "unavailable", "not disclosed"}:
                            if row.get("type") == "check" or label not in row.get("unavailable_reason_by_period", {}):
                                report("UNVERIFIABLE", addr, "Required or unexplained missing value")
                        if row.get("type") == "check":
                            counts["check_cells"] += 1
                            result = ev.value(short, rid, label)
                            if not is_number(result):
                                report("UNVERIFIABLE", addr, "Reconciliation cannot be evaluated; a missing input is not zero")
                            elif abs(result) > row.get("tolerance", 1e-9):
                                report("FAIL", addr, f"Reconciliation residual {result} exceeds explained tolerance")
                    except Exception as exc:
                        report("FAIL", addr, exc)
    for issue in ev.issues:
        report("UNVERIFIABLE", "arithmetic", issue)
    for ws in cached:
        for cells in ws.iter_rows():
            for cell in cells:
                if cell.data_type == "e":
                    report("FAIL", f"{ws.title}!{cell.coordinate}", f"Cached engine error: {cell.value}")
    for ws in formulas:
        for cells in ws.iter_rows():
            for cell in cells:
                if cell.data_type == "f" and "#REF!" in cell.value:
                    report("FAIL", f"{ws.title}!{cell.coordinate}", "Broken formula reference")
    if counts["check_cells"] == 0:
        report("FAIL", "checks", "No applicable reconciliation cells were evaluated")
    automated_status = ("FAIL" if any(f["status"] == "FAIL" for f in findings)
                        else "UNVERIFIABLE" if findings else "PASS")
    review_result = verify_agent_review(model, workbook, agent_review)
    findings.extend(review_result["findings"])
    status = ("FAIL" if any(f["status"] == "FAIL" for f in findings)
              else "UNVERIFIABLE" if findings else "PASS")
    return {"status": status, "verification_scope": "bundled_model_and_review_record",
            "automated_status": automated_status,
            "agent_review_status": review_result["status"], "review_binding": review_result["binding"],
            "visual_qa": "required separately",
            "counts": counts, "source_audit": {k: v for k, v in source_audit.items() if k != "findings"},
            "findings": findings}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("model", type=Path)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--workspace", type=Path, default=Path.cwd())
    parser.add_argument("--report-json", type=Path)
    parser.add_argument("--agent-review", type=Path, help="Independent reviewer Agent's check of originals and this final workbook")
    args = parser.parse_args()
    model = json.loads(args.model.read_text(encoding="utf-8"))
    if args.report_json:
        protected = {args.model.resolve(), args.workbook.resolve()}
        if args.agent_review:
            protected.add(args.agent_review.resolve())
        for source in model.get("sources", []):
            for relative in (source.get("path"), source.get("form_evidence", {}).get("path")):
                if relative:
                    protected.add((args.workspace / relative).resolve())
        if args.report_json.resolve() in protected:
            parser.error("Verification report cannot overwrite the model, workbook, review or original")
    review = json.loads(args.agent_review.read_text(encoding="utf-8")) if args.agent_review else None
    result = verify(model, args.workbook, args.workspace, agent_review=review)
    text = json.dumps(result, ensure_ascii=False, indent=2)
    if args.report_json:
        args.report_json.parent.mkdir(parents=True, exist_ok=True)
        args.report_json.write_text(text + "\n", encoding="utf-8")
    print(text)
    return {"PASS": 0, "FAIL": 2, "UNVERIFIABLE": 3}[result["status"]]


if __name__ == "__main__":
    raise SystemExit(main())
