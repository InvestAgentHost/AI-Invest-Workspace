#!/usr/bin/env python3
"""Build a NEW historical workbook using the approved shared style profile.

Default requires reviewed declarations and an independent raw-file audit.
--draft creates a conspicuously unverified preview. This never edits an existing
company workbook or any valuation extension. Engine recalculation is a later gate.
"""
from __future__ import annotations

import argparse
import json
import math
import unicodedata
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from model_contract import (DATA_START_ROW, cell_meta, contents_offset, is_number,
                            load_style, model_issues, sanitize_sheet_name,
                            ordered_sheets, sheet_role, tab_color, display_value)
from model_engine import FormulaPlan, Layout
from source_audit import audit_sources
from provenance import EvidenceIndex

# Keep build_registry import-compatible for existing initial-export helpers.
from model_contract import build_registry

PROFILE = load_style()


def fitted_height(text, width, minimum, font_size=10.5):
    """Conservative content fit, including explicit line breaks and CJK width."""
    units = lambda s: sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in str(s))
    capacity = max(4, (width - 2) * 10.5 / font_size)
    lines = sum(max(1, math.ceil(units(line) / capacity)) for line in str(text or "").split("\n"))
    heights = PROFILE["heights"]
    if lines == 1: return minimum
    if lines == 2: return max(minimum, heights["two_lines"])
    return max(minimum, heights["three_lines"] + (lines - 3) * 16)


class Builder:
    def __init__(self, model, draft=False):
        self.model, self.draft = model, draft
        self.evidence_index = EvidenceIndex(model)
        self.plan = FormulaPlan(model)
        self.layout = self.plan.layout
        self.registry, self.sheet_names = self.plan.registry, self.plan.sheet_names
        self.style, self.colors = PROFILE, PROFILE["colors"]
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.wb.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True)
        self.wb.properties.title = model.get("company", {}).get("name", "Historical research workbook")
        self.wb.properties.description = (
            "Historical 10-K/10-Q export. DRAFT — unverified."
            if draft else "Source-bound historical export; run engine recalculation and verification before release."
        )

    def font(self, role="body", color="text", bold=False, italic=False):
        return Font(name=self.style["font"], size=self.style["sizes"][role],
                    color=self.colors.get(color, color), bold=bold, italic=italic)

    def text(self, ws, r, c, value):
        cell = ws.cell(r, c, value)
        # External source text is always literal, never executable spreadsheet text.
        if isinstance(value, str): cell.data_type = "s"
        return cell

    def band(self, ws, row, title, last_col, role="section"):
        for ci in range(1, last_col + 1):
            cell = ws.cell(row, ci)
            cell.fill = PatternFill("solid", fgColor=self.colors["title"])
            cell.font = self.font(role, "white", bold=True)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
        self.text(ws, row, 1, title)
        ws.row_dimensions[row].height = fitted_height(
            title, self.style["widths"]["label"] + (last_col - 1) * self.style["widths"]["period"],
            self.style["heights"][role], self.style["sizes"][role])

    def base(self, name, last_col, title, scope):
        ws = self.wb.create_sheet(sanitize_sheet_name(name))
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = self.style["zoom"]
        self.band(ws, 1, ("DRAFT · " if self.draft else "") + title, last_col, "title")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
        cell = self.text(ws, 2, 1, scope)
        cell.font = self.font("note", "note")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[2].height = fitted_height(
            scope, self.style["widths"]["label"] + (last_col - 1) * self.style["widths"]["period"],
            self.style["heights"]["note"], self.style["sizes"]["note"])
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        return ws

    def statement(self, sheet):
        short = sheet["short"]
        co = self.model.get("company", {})
        scope = sheet.get("scope") or " | ".join(str(v) for v in (
            co.get("currency"), co.get("units") or co.get("scale"), co.get("reporting_framework"),
            f"As of {co['as_of_date']}" if co.get("as_of_date") else None) if v)
        ws = self.base(sheet["name"], self.layout.last_col, sheet.get("title", sheet["name"]), scope)
        ws.sheet_properties.tabColor = tab_color(sheet)
        widths, heights = self.style["widths"], self.style["heights"]
        ws.column_dimensions["A"].width = widths["label"]
        notes = any(row.get("note") for block in sheet.get("blocks", []) for row in block.get("rows", []))
        ws.column_dimensions["B"].width = widths["note"] if notes else widths["gutter"]
        ws.column_dimensions["B"].hidden = bool(notes)
        ws.column_dimensions["B"].outlineLevel = 1 if notes else 0
        ws.column_dimensions["B"].collapsed = bool(notes)
        ws.row_dimensions[3].height = heights["header"]
        ws.row_dimensions[4].height = heights["date"]
        ws.row_dimensions[5].height = heights["note"]
        self.text(ws, 3, 1, "Metric / 科目").font = self.font("header", bold=True)
        self.text(ws, 5, 1, "Units / period basis: see section notes").font = self.font("note", "note")
        for info in self.layout.entries:
            label, ci = info["label"], self.layout.col_of[info["label"]]
            ws.column_dimensions[get_column_letter(ci)].width = widths["period"]
            color = "annual_header" if label in self.layout.annual_labels else "quarter_header"
            cell = self.text(ws, 3, ci, label)
            cell.font = self.font("header", "white", bold=True)
            cell.fill = PatternFill("solid", fgColor=self.colors[color])
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if info.get("end"):
                cell = ws.cell(4, ci, date.fromisoformat(info["end"]))
                cell.number_format = "yyyy-mm-dd"
                cell.font = self.font("date", "note")
                cell.fill = PatternFill("solid", fgColor=self.colors["date_fill"])
                cell.alignment = Alignment(horizontal="center", vertical="center")
            grain = info.get("grain", "annual" if label in self.layout.annual_labels else "discrete")
            self.text(ws, 5, ci, grain).font = self.font("note", "note")
        for ci in self.layout.scrap_cols:
            ws.column_dimensions[get_column_letter(ci)].width = widths["separator"]
        ws.freeze_panes = "C6"
        ws.print_title_rows = "1:5"
        if sheet.get("contents"):
            for i, block in enumerate(sheet.get("blocks", []), DATA_START_ROW):
                cell = self.text(ws, i, 1, block["title"])
                target = self.plan.block_rows[short][i - DATA_START_ROW]
                safe_name = ws.title.replace("'", "''")
                cell.hyperlink = f"#'{safe_name}'!A{target}"
                cell.font = self.font("note", "link")
                ws.row_dimensions[i].height = fitted_height(block["title"], widths["label"], heights["note"], 9)
        for bi, block in enumerate(sheet.get("blocks", [])):
            r = self.plan.block_rows[short][bi]
            self.band(ws, r, block["title"], self.layout.last_col)
            r += 1
            if block.get("note"):
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=self.layout.last_col)
                cell = self.text(ws, r, 1, block["note"])
                cell.font = self.font("note", "note")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
                ws.row_dimensions[r].height = fitted_height(
                    block["note"], widths["label"] + (self.layout.last_col - 1) * widths["period"], heights["note"], 9)
                r += 1
            for row in block.get("rows", []):
                kind = row.get("type", "input")
                if kind == "spacer":
                    ws.row_dimensions[r].height = heights["gap"]
                    r += 1
                    continue
                if kind == "note":
                    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=self.layout.last_col)
                    cell = self.text(ws, r, 1, row.get("label", ""))
                    cell.font = self.font("note", "note")
                    cell.alignment = Alignment(wrap_text=True, vertical="center")
                    ws.row_dimensions[r].height = fitted_height(
                        row.get("label", ""), widths["label"] + (self.layout.last_col - 1) * widths["period"], heights["note"], 9)
                    r += 1
                    continue
                ratio = row.get("fmt") in {"pct", "x"} or kind.startswith("growth_")
                total = row.get("bold", False) or kind == "subtotal"
                minimum = heights["total"] if total else heights["ratio"] if ratio else heights["body"]
                label = self.text(ws, r, 1, row.get("label", row["id"]))
                label.font = self.font(bold=total, italic=ratio)
                label.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True,
                                            indent=int(row.get("indent", 0)))
                h = fitted_height(label.value, widths["label"] - int(row.get("indent", 0)) * 2, minimum)
                ws.row_dimensions[r].height = max(h, row.get("row_height", 0))
                if row.get("note"):
                    cell = self.text(ws, r, 2, row["note"])
                    cell.font = self.font("note", "note")
                    cell.alignment = Alignment(wrap_text=True, vertical="center")
                    # Notes are expandable, so their content also gets adequate height.
                    ws.row_dimensions[r].height = max(ws.row_dimensions[r].height,
                        fitted_height(row["note"], widths["note"], minimum, 9))
                for period in self.layout.labels:
                    ci = self.layout.col_of[period]
                    value = self.plan.cell_value(short, row, period)
                    cell = ws.cell(r, ci, value)
                    formula = isinstance(value, str) and value.startswith("=")
                    if isinstance(value, str) and not formula: cell.data_type = "s"
                    color = "input" if is_number(value) else "link" if formula and "!" in value else "text"
                    cell.font = self.font(color=color, bold=total, italic=ratio)
                    cell.number_format = self.style["formats"][row.get("fmt", "pct" if kind.startswith("growth_") else "num")]
                    cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=False)
                    if kind == "check" and formula:
                        tolerance = row.get("tolerance", 1e-9)
                        ws.conditional_formatting.add(cell.coordinate, CellIsRule(
                            operator="notBetween", formula=[str(-tolerance), str(tolerance)],
                            fill=PatternFill("solid", fgColor=self.colors["error_fill"]),
                            font=self.font(color="error_text", bold=True)))
                if total:
                    for ci in [1, *self.layout.col_of.values()]:
                        cell = ws.cell(r, ci)
                        cell.fill = PatternFill("solid", fgColor=self.colors["subtotal"])
                        cell.border = Border(top=Side(style="thin", color=self.colors["border"]))
                r += 1
            ws.row_dimensions[r].height = heights["gap"]
        ws.print_options.horizontalCentered = False

    def table(self, sheet):
        rows = sheet.get("table", [])
        evidence = {(e["row"], e["column"]): e for e in sheet.get("cell_evidence", [])}
        last_col = max(2, max((len(r) for r in rows), default=2))
        ws = self.base(sheet["name"], last_col, sheet.get("title", sheet["name"]), sheet.get("scope", ""))
        ws.sheet_properties.tabColor = tab_color(sheet)
        is_contents = sheet_role(sheet) == 'contents'
        for ci in range(1, last_col + 1):
            ws.column_dimensions[get_column_letter(ci)].width = 60 if ci == 1 else 100 if is_contents else 34
        for ri, record in enumerate(rows, 4):
            h = self.style["heights"]["contents" if is_contents else "body"]
            for ci, value in enumerate(record, 1):
                cell = self.text(ws, ri, ci, display_value(value))
                header = ri == 4
                cell.font = self.font("header" if header else "note",
                                      "white" if header else "input" if is_number(value) else "text", bold=header)
                if is_number(value):
                    ref = evidence.get((ri - 4, ci - 1), {})
                    if ref.get("table"):
                        ref = self.evidence_index.resolve({"id": ref["id"]}, ref["period_label"], ref)
                    cell.number_format = self.style["formats"][ref.get("fmt", "num")]
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if header: cell.fill = PatternFill("solid", fgColor=self.colors["annual_header"])
                if not is_contents:h = max(h, fitted_height(value, 60 if ci == 1 else 34, h, 10 if header else 9))
            ws.row_dimensions[ri].height = h
        ws.freeze_panes = "B5"

    def run(self, output, workspace=None, overwrite=False):
        output = Path(output)
        if output.exists() and not overwrite:
            raise FileExistsError("Initial export will not overwrite a workbook; use a new path")
        if not self.draft:
            issues = model_issues(self.model)
            if issues: raise ValueError("Declaration gate failed:\n" + "\n".join(issues))
            audit = audit_sources(self.model, workspace or Path.cwd())
            if audit["status"] != "PASS":
                raise ValueError("Original-source audit failed:\n" + json.dumps(audit, ensure_ascii=False, indent=2))
        for sheet in ordered_sheets(self.model["sheets"]):
            if sheet.get("kind") == "statement": self.statement(sheet)
            elif sheet.get("kind") == "table": self.table(sheet)
            else: raise ValueError(f"Unknown sheet kind: {sheet.get('kind')}")
        output.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("model", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--workspace", type=Path, default=Path.cwd())
    parser.add_argument("--draft", action="store_true", help="Unverified visual preview; never a release")
    # No generic --overwrite: existing manually edited workbooks use scoped editing.
    args = parser.parse_args()
    model = json.loads(args.model.read_text(encoding="utf-8"))
    Builder(model, draft=args.draft).run(args.output, workspace=args.workspace)
    print(json.dumps({"status": "DRAFT" if args.draft else "BUILT_PENDING_ENGINE_AND_AGENT_REVIEW",
                      "output": str(args.output)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
