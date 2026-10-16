"""Independent recalculation of KO_model.xlsx formulas straight from the file; reports Excel-style errors."""
import re, sys
from openpyxl import load_workbook
wb = load_workbook(sys.argv[1])
REF = re.compile(r"(?:'([^']+)'!)?\$?([A-Z]{1,3})\$?(\d+)")
cache, errs = {}, []

def val(ws, coord, stack=()):
    key = (ws.title, coord)
    if key in cache: return cache[key]
    if key in stack: raise RuntimeError("circular " + str(key))
    v = ws[coord].value
    if isinstance(v, str) and v.startswith("=") and not v.startswith("=HYPERLINK"):
        v = ev(ws, v[1:], stack + (key,), coord)
    cache[key] = v
    return v

def ev(ws, f, stack, coord):
    def sub(m):
        sheet = wb[m.group(1)] if m.group(1) else ws
        x = val(sheet, m.group(2) + m.group(3), stack)
        if isinstance(x, str) and x == "": return "0.0"
        if x is None: return "BLANK"
        if isinstance(x, str): return "STR"
        return repr(float(x))
    py = REF.sub(sub, f)
    py = py.replace("N(", "_N(")
    def _N(x): return 0 if x is None or isinstance(x, str) else x
    if "STR" in py:
        errs.append((ws.title, coord, "#VALUE!", f)); return None
    try:
        if "BLANK" in py:
            # blank in arithmetic is 0 in Excel; blank denominator => #DIV/0!
            py2 = py.replace("BLANK", "0.0")
            r = eval(py2.replace('IF(', '_IF(').replace('"",', '"",'), {"_N": lambda x: x, "_IF": lambda c, a, b: a if c else b}, {})
        else:
            r = eval(py.replace('IF(', '_IF('), {"_N": lambda x: x, "_IF": lambda c, a, b: a if c else b}, {})
        return r
    except ZeroDivisionError:
        errs.append((ws.title, coord, "#DIV/0!", f)); return None
    except Exception as e:
        errs.append((ws.title, coord, "#ERR " + type(e).__name__, f)); return None

# patch IF to be lazy for the growth pattern: IF(N(prev)=0,"",cur/prev-1) -- handle by regex first
GROW = re.compile(r'^IF\(N\((.+?)\)=0,"",(.+)\)$')
orig_ev = ev
def ev(ws, f, stack, coord):
    m = GROW.match(f)
    if m:
        prev = m.group(1)
        pv = orig_ev(ws, prev, stack, coord) if REF.fullmatch(prev) or True else None
        if pv in (None, 0, "") : return ""
        return orig_ev(ws, m.group(2), stack, coord)
    return orig_ev(ws, f, stack, coord)

n = 0; bad_checks = []
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("=") and not c.value.startswith("=HYPERLINK"):
                n += 1
                v = val(ws, c.coordinate)
                lab = ws.cell(row=c.row, column=1).value or ""
                if "Check" in lab and isinstance(v, (int, float)) and abs(v) > 0.5 and "Info" not in lab and "rounding" not in lab:
                    bad_checks.append((ws.title, c.coordinate, lab[:60], v))
print("formula cells evaluated:", n)
print("error cells:", len(errs))
by = {}
for s, c, e, f in errs: by.setdefault((s, e), []).append(c)
for k, v in by.items(): print(" ", k, len(v), v[:8])
print("check rows off:", bad_checks[:10])
