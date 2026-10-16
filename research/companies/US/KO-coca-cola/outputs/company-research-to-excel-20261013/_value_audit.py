import json, re, sys
from openpyxl import load_workbook
ROOT = "/Users/tccc/Desktop/AI Invest/Workspace/"
W = ROOT + "research/companies/US/KO-coca-cola/outputs/company-research-to-excel-20261013/"
RAW = json.load(open(W + "statements_raw.json"))
wb = load_workbook(W + "KO_model.xlsx")
QKEY = {"3Q2024": "10Q-Q3/FY2024", "1Q2025": "10Q-Q1/FY2025", "2Q2025": "10Q-Q2/FY2025", "3Q2025": "10Q-Q3/FY2025", "1Q2026": "10Q-Q1/FY2026", "2Q2026": "10Q-Q2/FY2026"}
cache = {}
def lines(key):
    if key not in cache:
        cache[key] = open(ROOT + RAW[key]["md_path"], encoding="utf-8").read().split("\n")
    return cache[key]
def src_for(label, stmt):
    if label in QKEY: return QKEY[label]
    y = int(label[2:]); off = 1 if stmt == "BS" else 2
    return f"10K-FY{min(y + off, 2025)}"
def fmt(v):
    s = f"{abs(v):,.2f}" if isinstance(v, float) and v != int(v) else f"{abs(int(v)):,}"
    return s
ok = bad = 0; samples = []
for name, stmt in (("Income Statement", "IS"), ("Balance Sheet", "BS"), ("Cash Flow", "CF")):
    ws = wb[name]
    hdr = {c.column: c.value for c in ws[3] if c.value}
    for row in ws.iter_rows(min_row=6):
        label = row[0].value
        note = ws.cell(row=row[0].row, column=2).value or ""
        if "Filed caption: " in note:
            label = note.split("Filed caption: ", 1)[1].strip()
        if not label or label.startswith(("Check", "Balance Check", "Cash Walk", "  ", "Days", "Cash Conversion")): continue
        for c in row[2:]:
            p = hdr.get(c.column)
            if not p or p == "·" or isinstance(c.value, str) or c.value is None: continue
            key = src_for(p, stmt)
            L = lines(key)
            first = re.sub(r"\s+\d+$", "", label.strip())[:22].lower()
            found = False
            for l in L:
                if l.startswith("|") and l[1:].strip().lower().startswith(first[:18]):
                    nums = re.findall(r"\(?\d[\d,]*\.?\d*\)?", re.sub(r"\s*\|\s*\)", ")", l))
                    if fmt(c.value) in [n.strip("()") for n in nums]:
                        found = True; break
            if not found:
                # fallback: same number anywhere in the statement table of the source filing (caption renamed)
                tb = RAW[key][stmt]
                target = c.value
                yrs = [int(m) for h in tb["headers"][1:] for m in re.findall(r"20\d\d", h)]
                for lab, vals in tb["rows"]:
                    if any(isinstance(x, (int, float)) and x == target for x in vals) and lab and lab.lower()[:8] in (label.lower()[:8], "other as", "other li", "common s", "income f", "income b"):
                        found = True; break
            if found: ok += 1
            else:
                bad += 1
                if len(samples) < 25: samples.append((name, label[:50], p, c.value, key))
print("inputs found on source line:", ok, " not matched:", bad)
for s in samples: print(" ", s)
