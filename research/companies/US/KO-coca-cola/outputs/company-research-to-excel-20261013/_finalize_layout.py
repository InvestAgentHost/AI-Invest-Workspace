"""Formatting-only post-step for KO_model.xlsx (no values/formulas are changed):
widen label column on statement sheets, wrap long text on Cover/Sources, set content-fitted row heights."""
import math, sys
from openpyxl import load_workbook
from openpyxl.styles import Alignment
src, dst = sys.argv[1], sys.argv[2]
wb = load_workbook(src)
for ws in wb.worksheets:
    if ws.title in ("Cover", "Sources"):
        widths = {c: (ws.column_dimensions[c].width or 10) for c in "ABCDE"}
        for row in ws.iter_rows(min_row=4):
            need = 1
            for c in row:
                if isinstance(c.value, str) and c.value:
                    w = widths.get(c.column_letter, 20)
                    c.alignment = Alignment(wrap_text=True, vertical="top")
                    need = max(need, math.ceil(len(c.value) * 1.05 / max(w - 2, 8)))
            if need > 1:
                ws.row_dimensions[row[0].row].height = 13.5 * need
    else:
        ws.column_dimensions["A"].width = 74
wb.save(dst)
print("finalized", dst)
